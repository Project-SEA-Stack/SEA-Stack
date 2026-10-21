#!/usr/bin/env python3
"""
Normalize the IEA OES Task 10 sphere heave RAO data at steepness S = 0.01.

Raw input (shared with sphere_rao_sweep, not duplicated):
  ../sphere_rao_sweep/raw/RAO_dat.xlsx
    Sheet noPTO_S=0.01 -> free body
    Sheet PTO_S=0.01   -> optimum PTO damping
    Columns: Wave Period (s), InWave-HOTINT, InWave (Lin), Marin (NLin),
             NREL (NLin), ProteusDS (Lin)
    Values are heave RAO in m/m.

Steepness is S = H / (g T^2) (Wendt et al. 2019, eq. 5), so H = S g T^2 and
the wave amplitude is a = H / 2.

Output (normalized/):
  <source>/heave_rao_free.txt and <source>/heave_rao_pto.txt:
  omega (rad/s), heave RAO amplitude (m/m).
"""

import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

import numpy as np
from normalize_utils import write_rao

HERE = os.path.dirname(os.path.abspath(__file__))

SOURCES = {
    'InWave-HOTINT': 'inwave_hotint',
    'InWave (Lin)':  'inwave',
    'Marin (NLin)':  'marin',
    'NREL (NLin)':   'nrel',
    'ProteusDS (Lin)': 'proteusds',
}
SHEETS = {'noPTO_S=0.01': 'free', 'PTO_S=0.01': 'pto'}


def main():
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl required (pip install openpyxl)")
        sys.exit(1)

    xlsx = os.path.join(HERE, '..', 'sphere_rao_sweep', 'raw', 'RAO_dat.xlsx')
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    for sheet, config in SHEETS.items():
        ws = wb[sheet]
        header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        periods, source_data = [], {v: [] for v in SOURCES.values()}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                break
            periods.append(float(row[0]))
            for col_name, dir_name in SOURCES.items():
                source_data[dir_name].append(float(row[header.index(col_name)]))

        omega = 2.0 * np.pi / np.array(periods)
        for dir_name, rao_vals in source_data.items():
            write_rao(
                os.path.join(HERE, 'normalized', dir_name, f'heave_rao_{config}.txt'),
                np.column_stack([omega, np.array(rao_vals)]),
                source=dir_name, model='sphere_rao_sweep_nl',
                quantity=f'heave_rao_{config}',
                units='rad/s, m/m',
                columns='Omega(rad/s) HeaveRAO(m/m)',
                extra_comments=[f'sheet: {sheet} (S = H/(g T^2) = 0.01)'],
            )
    wb.close()
    print(f"Wrote {len(SOURCES) * len(SHEETS)} normalized sphere RAO files")


if __name__ == '__main__':
    main()
